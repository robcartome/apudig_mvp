import io
from collections import defaultdict
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import BankMovement, CashSafeMovement


def _safe_text(value):
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def build_cash_safe_report(movements, *, store_name, date_from=None, date_to=None):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    detail = workbook.create_sheet("Movimientos")
    header_fill = PatternFill("solid", fgColor="1A7F64")
    header_font = Font(bold=True, color="FFFFFF")

    period = "Todos los movimientos"
    if date_from or date_to:
        period = f"Del {date_from:%d/%m/%Y}" if date_from else "Desde el inicio"
        period += f" al {date_to:%d/%m/%Y}" if date_to else " en adelante"
    summary.append(["Reporte de fondos y caja fuerte"])
    summary["A1"].font = Font(bold=True, size=16, color="123B34")
    summary.append(["Sucursal", _safe_text(store_name)])
    summary.append(["Periodo", period])
    summary.append(["Generado", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    summary.append([])
    summary.append(["Fondo", "Moneda", "Movimientos", "Ingresos confirmados", "Egresos confirmados", "Variación neta"])
    for cell in summary[6]:
        cell.fill = header_fill
        cell.font = header_font

    totals = defaultdict(lambda: {
        "count": 0, "incoming": Decimal("0"), "outgoing": Decimal("0")
    })
    rows = list(movements)
    for movement in rows:
        key = (movement.fund.name, movement.currency)
        totals[key]["count"] += 1
        if movement.status in {
            CashSafeMovement.Status.POSTED, CashSafeMovement.Status.REVERSED,
        }:
            target = "incoming" if movement.direction == CashSafeMovement.Direction.IN else "outgoing"
            totals[key][target] += movement.amount
    for (fund, currency), item in sorted(totals.items()):
        summary.append([
            _safe_text(fund), currency, item["count"], float(item["incoming"]),
            float(item["outgoing"]), float(item["incoming"] - item["outgoing"]),
        ])

    headers = [
        "Fecha", "Documento", "Fondo", "Dirección", "Tipo", "Estado",
        "Moneda", "Importe", "Responsable", "Destino", "Caja POS destino", "Referencia",
        "Registrado por", "Autorizado por", "Fecha autorización", "Proveedor",
        "Documento de compra", "Categoría", "N.º comprobante", "Observación",
    ]
    detail.append(headers)
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for movement in rows:
        detail.append([
            timezone.localtime(movement.created_at).replace(tzinfo=None),
            movement.movement_number, _safe_text(movement.fund.name),
            movement.get_direction_display(), movement.get_movement_type_display(),
            movement.get_status_display(), movement.currency, float(movement.amount),
            _safe_text(movement.recipient_name), _safe_text(movement.destination),
            _safe_text(movement.destination_register),
            _safe_text(movement.reference),
            _safe_text(movement.created_by.display_name if movement.created_by else "Sistema"),
            _safe_text(movement.authorized_by.display_name if movement.authorized_by else ""),
            timezone.localtime(movement.authorized_at).replace(tzinfo=None) if movement.authorized_at else None,
            _safe_text(movement.supplier), _safe_text(movement.purchase_document),
            _safe_text(movement.expense_category),
            _safe_text(movement.supporting_document_number), _safe_text(movement.notes),
        ])
    for worksheet in (summary, detail):
        worksheet.freeze_panes = "A7" if worksheet is summary else "A2"
        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions if worksheet is detail else f"A6:F{worksheet.max_row}"
        for column in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column)
            width = max(len(str(cell.value or "")) for cell in worksheet[letter]) + 3
            worksheet.column_dimensions[letter].width = min(max(width, 12), 42)
    for row in detail.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "dd/mm/yyyy hh:mm"
    for row in detail.iter_rows(min_row=2, min_col=15, max_col=15):
        row[0].number_format = "dd/mm/yyyy hh:mm"
    for row in summary.iter_rows(min_row=7, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"
    for row in detail.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = "#,##0.00"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_bank_movement_report(movements, *, company_name, date_from=None, date_to=None):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    detail = workbook.create_sheet("Movimientos bancarios")
    header_fill = PatternFill("solid", fgColor="1A7F64")
    header_font = Font(bold=True, color="FFFFFF")
    period = "Todos los movimientos"
    if date_from or date_to:
        period = f"Del {date_from:%d/%m/%Y}" if date_from else "Desde el inicio"
        period += f" al {date_to:%d/%m/%Y}" if date_to else " en adelante"
    summary.append(["Reporte de cuentas y movimientos bancarios"])
    summary["A1"].font = Font(bold=True, size=16, color="123B34")
    summary.append(["Empresa", _safe_text(company_name)])
    summary.append(["Periodo", period])
    summary.append(["Generado", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    summary.append([])
    summary.append(["Cuenta", "Moneda", "Movimientos", "Ingresos", "Egresos", "Variación neta"])
    for cell in summary[6]:
        cell.fill = header_fill
        cell.font = header_font
    rows = list(movements)
    totals = defaultdict(lambda: {"count": 0, "incoming": Decimal("0"), "outgoing": Decimal("0")})
    for movement in rows:
        key = (str(movement.account), movement.account.currency)
        totals[key]["count"] += 1
        if movement.status in {BankMovement.Status.POSTED, BankMovement.Status.REVERSED}:
            target = "incoming" if movement.direction == BankMovement.Direction.IN else "outgoing"
            totals[key][target] += movement.amount
    for (account, currency), item in sorted(totals.items()):
        summary.append([
            _safe_text(account), currency, item["count"], float(item["incoming"]),
            float(item["outgoing"]), float(item["incoming"] - item["outgoing"]),
        ])
    detail.append([
        "Fecha", "Documento", "Cuenta", "Dirección", "Tipo", "Estado",
        "Moneda", "Importe", "Referencia", "Descripción", "Cuenta destino",
        "Registrado por", "Reversión / motivo",
    ])
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for movement in rows:
        detail.append([
            movement.operation_date, movement.movement_number, _safe_text(movement.account),
            movement.get_direction_display(), movement.get_movement_type_display(),
            movement.get_status_display(), movement.account.currency, float(movement.amount),
            _safe_text(movement.reference), _safe_text(movement.description),
            _safe_text(movement.destination_account),
            _safe_text(movement.created_by.display_name if movement.created_by else "Sistema"),
            _safe_text(movement.reversal_reason),
        ])
    for worksheet in (summary, detail):
        worksheet.freeze_panes = "A7" if worksheet is summary else "A2"
        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions if worksheet is detail else f"A6:F{worksheet.max_row}"
        for column in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column)
            width = max(len(str(cell.value or "")) for cell in worksheet[letter]) + 3
            worksheet.column_dimensions[letter].width = min(max(width, 12), 42)
    for row in summary.iter_rows(min_row=7, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"
    for row in detail.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "dd/mm/yyyy"
    for row in detail.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].number_format = "#,##0.00"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
