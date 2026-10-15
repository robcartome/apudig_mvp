import csv
import hashlib
import io
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from django.db.models import Max, Q, Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.companies.models import Company, Store
from apps.core.models import AuditLog
from apps.users.permissions import user_store_role
from .models import (
    BankAccount,
    BankMovement,
    BankReconciliationMatch,
    BankStatement,
    BankStatementLine,
    CashSafeDeposit,
    CashSafeMovement,
    CashSafeMovementAttachment,
    TreasuryFund,
)


class TreasuryDomainError(Exception):
    pass


def add_cash_safe_attachment(*, movement, uploaded_file, user):
    attachment = CashSafeMovementAttachment(
        movement=movement, file=uploaded_file,
        original_name=Path(uploaded_file.name).name[:255],
        content_type=uploaded_file.content_type or "application/octet-stream",
        size=uploaded_file.size, uploaded_by=user,
    )
    attachment.save()
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_ATTACHMENT_ADD",
        entity="CashSafeMovement", entity_id=str(movement.pk),
        meta_data={"attachment_id": str(attachment.pk), "name": attachment.original_name},
    )
    return attachment


def get_or_create_default_fund(*, company_id, store_id):
    fund = TreasuryFund.objects.filter(
        company_id=company_id, store_id=store_id, is_default=True
    ).first()
    if fund:
        return fund
    return TreasuryFund.objects.create(
        company_id=company_id, store_id=store_id, code="PRINCIPAL",
        name="Caja fuerte principal", fund_type=TreasuryFund.FundType.SAFE,
        is_default=True, active=True,
    )


def cash_safe_balance(*, company_id, store_id, currency, fund_id=None, at=None):
    movements = CashSafeMovement.objects.filter(
        company_id=company_id, store_id=store_id, currency=currency,
        status__in=(CashSafeMovement.Status.POSTED, CashSafeMovement.Status.REVERSED),
    )
    if fund_id:
        movements = movements.filter(fund_id=fund_id)
    if at:
        movements = movements.filter(created_at__lte=at)
    totals = movements.aggregate(
        incoming=Sum("amount", filter=Q(direction=CashSafeMovement.Direction.IN)),
        outgoing=Sum("amount", filter=Q(direction=CashSafeMovement.Direction.OUT)),
    )
    return (totals["incoming"] or Decimal("0.00")) - (
        totals["outgoing"] or Decimal("0.00")
    )


def _next_cash_safe_number(company_id, occurred_at=None):
    Company.objects.select_for_update().get(pk=company_id)
    local_date = (
        timezone.localtime(occurred_at).date()
        if occurred_at and timezone.is_aware(occurred_at)
        else (occurred_at.date() if occurred_at else timezone.localdate())
    )
    prefix = f"CF-{local_date:%Y%m}-"
    sequence = (
        CashSafeMovement.objects.filter(
            company_id=company_id, movement_number__startswith=prefix
        ).aggregate(value=Max("sequence_number"))["value"] or 0
    ) + 1
    return sequence, f"{prefix}{sequence:06d}"


def _next_bank_movement_number(company_id, occurred_at=None):
    Company.objects.select_for_update().get(pk=company_id)
    local_date = occurred_at or timezone.localdate()
    if hasattr(local_date, "date"):
        local_date = local_date.date()
    prefix = f"BM-{local_date:%Y%m}-"
    sequence = (
        BankMovement.objects.filter(
            company_id=company_id, movement_number__startswith=prefix
        ).aggregate(value=Max("sequence_number"))["value"] or 0
    ) + 1
    return sequence, f"{prefix}{sequence:06d}"


def bank_account_balance(account):
    totals = account.movements.filter(
        status__in=(BankMovement.Status.POSTED, BankMovement.Status.REVERSED)
    ).aggregate(
        incoming=Sum("amount", filter=Q(direction=BankMovement.Direction.IN)),
        outgoing=Sum("amount", filter=Q(direction=BankMovement.Direction.OUT)),
    )
    return account.opening_balance + (totals["incoming"] or Decimal("0.00")) - (
        totals["outgoing"] or Decimal("0.00")
    )


def bank_account_balance_at(account, date_to):
    totals = account.movements.filter(
        status__in=(BankMovement.Status.POSTED, BankMovement.Status.REVERSED),
        operation_date__lte=date_to,
    ).aggregate(
        incoming=Sum("amount", filter=Q(direction=BankMovement.Direction.IN)),
        outgoing=Sum("amount", filter=Q(direction=BankMovement.Direction.OUT)),
    )
    return account.opening_balance + (totals["incoming"] or Decimal("0.00")) - (
        totals["outgoing"] or Decimal("0.00")
    )


def _normalized_header(value):
    plain = unicodedata.normalize("NFKD", value or "").encode("ascii", "ignore").decode()
    return "_".join(plain.strip().lower().split())


def _parse_statement_decimal(value):
    raw = str(value or "").strip().replace("S/", "").replace("$", "").replace(" ", "")
    if not raw:
        return Decimal("0.00")
    negative = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()")
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        tail = raw.rsplit(",", 1)[1]
        raw = raw.replace(",", ".") if len(tail) <= 2 else raw.replace(",", "")
    try:
        amount = Decimal(raw)
    except InvalidOperation as exc:
        raise TreasuryDomainError(f"Importe inválido en el extracto: {value}") from exc
    return -amount if negative else amount


def _parse_statement_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime((value or "").strip(), date_format).date()
        except ValueError:
            continue
    raise TreasuryDomainError(f"Fecha inválida en el extracto: {value}")


def _read_statement_rows(uploaded_file):
    extension = Path(uploaded_file.name).suffix.lower()
    if extension == ".xlsx":
        try:
            workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
            worksheet = workbook.active
            iterator = worksheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator)]
        except (StopIteration, ValueError, OSError) as exc:
            raise TreasuryDomainError("El archivo Excel no contiene una hoja válida.") from exc
        return headers, [dict(zip(headers, row)) for row in iterator]
    if extension != ".csv":
        raise TreasuryDomainError("El extracto debe ser un archivo CSV o XLSX.")
    try:
        content = uploaded_file.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        uploaded_file.seek(0)
        content = uploaded_file.read().decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(content), dialect=dialect)
    return reader.fieldnames or [], list(reader)


@transaction.atomic
def import_bank_statement(*, company_id, user, uploaded_file, **data):
    account = BankAccount.objects.get(pk=data["account"].pk, company_id=company_id)
    fieldnames, statement_rows = _read_statement_rows(uploaded_file)
    if not fieldnames:
        raise TreasuryDomainError("El archivo CSV no contiene encabezados.")
    header_map = {_normalized_header(name): name for name in fieldnames}
    date_key = header_map.get("fecha")
    description_key = header_map.get("descripcion") or header_map.get("concepto")
    reference_key = header_map.get("referencia") or header_map.get("documento")
    debit_key = header_map.get("cargo") or header_map.get("debito")
    credit_key = header_map.get("abono") or header_map.get("credito")
    amount_key = header_map.get("monto") or header_map.get("importe")
    if not date_key or not description_key or not ((debit_key and credit_key) or amount_key):
        raise TreasuryDomainError(
            "El CSV debe incluir fecha, descripción y cargo/abono, o una columna monto."
        )

    parsed_rows = []
    for row_number, row in enumerate(statement_rows, start=2):
        if not any(str(value or "").strip() for value in row.values()):
            continue
        transaction_date = _parse_statement_date(row.get(date_key))
        if not data["date_from"] <= transaction_date <= data["date_to"]:
            raise TreasuryDomainError(
                f"La fecha de la fila {row_number} está fuera del periodo seleccionado."
            )
        if amount_key:
            signed_amount = _parse_statement_decimal(row.get(amount_key))
            if not signed_amount:
                raise TreasuryDomainError(f"La fila {row_number} tiene importe cero.")
            direction = (
                BankMovement.Direction.IN if signed_amount > 0 else BankMovement.Direction.OUT
            )
            amount = abs(signed_amount)
        else:
            debit = abs(_parse_statement_decimal(row.get(debit_key)))
            credit = abs(_parse_statement_decimal(row.get(credit_key)))
            if bool(debit) == bool(credit):
                raise TreasuryDomainError(
                    f"La fila {row_number} debe tener solo cargo o solo abono."
                )
            direction = BankMovement.Direction.OUT if debit else BankMovement.Direction.IN
            amount = debit or credit
        description = (row.get(description_key) or "").strip()
        reference = (row.get(reference_key) or "").strip() if reference_key else ""
        fingerprint_source = (
            f"{row_number}|{transaction_date}|{direction}|{amount}|{reference}|{description}"
        )
        parsed_rows.append({
            "transaction_date": transaction_date,
            "direction": direction,
            "amount": amount,
            "reference": reference,
            "description": description,
            "fingerprint": hashlib.sha256(fingerprint_source.encode()).hexdigest(),
        })
    if not parsed_rows:
        raise TreasuryDomainError("El extracto no contiene movimientos para importar.")

    statement = BankStatement(
        company_id=company_id,
        account=account,
        date_from=data["date_from"],
        date_to=data["date_to"],
        opening_balance=data["opening_balance"],
        closing_balance=data["closing_balance"],
        source_filename=uploaded_file.name[:255],
        created_by=user,
    )
    statement.full_clean()
    statement.save()
    BankStatementLine.objects.bulk_create(
        [BankStatementLine(statement=statement, **row) for row in parsed_rows]
    )
    AuditLog.objects.create(
        user=user, action="BANK_STATEMENT_IMPORT", entity="BankStatement",
        entity_id=str(statement.pk), meta_data={"rows": len(parsed_rows)},
    )
    return statement


@transaction.atomic
def match_bank_statement_line(*, line_id, movement_id, company_id, user):
    line = BankStatementLine.objects.select_related("statement").select_for_update().get(
        pk=line_id, statement__company_id=company_id
    )
    if line.statement.status != BankStatement.Status.DRAFT:
        raise TreasuryDomainError("La conciliación ya está cerrada.")
    movement = BankMovement.objects.select_for_update().get(
        pk=movement_id, company_id=company_id, account=line.statement.account,
        status=BankMovement.Status.POSTED,
    )
    if movement.direction != line.direction:
        raise TreasuryDomainError("El movimiento debe coincidir con la dirección del extracto.")
    if not line.statement.date_from <= movement.operation_date <= line.statement.date_to:
        raise TreasuryDomainError("El movimiento está fuera del periodo de conciliación.")
    if BankReconciliationMatch.objects.filter(bank_movement=movement).exclude(statement_line=line).exists():
        raise TreasuryDomainError("El movimiento del sistema ya está conciliado.")
    matched_total = line.reconciliation_matches.aggregate(
        total=Sum("bank_movement__amount")
    )["total"] or Decimal("0.00")
    if matched_total + movement.amount > line.amount:
        raise TreasuryDomainError("La suma seleccionada supera el importe del extracto.")
    match, created = BankReconciliationMatch.objects.get_or_create(
        statement_line=line, bank_movement=movement,
        defaults={"matched_by": user},
    )
    return match


@transaction.atomic
def unmatch_bank_statement_line(*, line_id, company_id, movement_id=None):
    line = BankStatementLine.objects.select_related("statement").get(
        pk=line_id, statement__company_id=company_id
    )
    if line.statement.status != BankStatement.Status.DRAFT:
        raise TreasuryDomainError("La conciliación ya está cerrada.")
    matches = BankReconciliationMatch.objects.filter(statement_line=line)
    if movement_id:
        matches = matches.filter(bank_movement_id=movement_id)
    matches.delete()


@transaction.atomic
def close_bank_statement(*, statement_id, company_id, user):
    statement = BankStatement.objects.select_for_update().get(
        pk=statement_id, company_id=company_id
    )
    if statement.status != BankStatement.Status.DRAFT:
        raise TreasuryDomainError("La conciliación ya está cerrada.")
    for line in statement.lines.prefetch_related("reconciliation_matches__bank_movement"):
        matched_total = sum(
            (match.bank_movement.amount for match in line.reconciliation_matches.all()),
            Decimal("0.00"),
        )
        if matched_total != line.amount:
            raise TreasuryDomainError(
                "Aún existen movimientos bancarios sin conciliar completamente."
            )
    totals = statement.lines.aggregate(
        incoming=Sum("amount", filter=Q(direction=BankMovement.Direction.IN)),
        outgoing=Sum("amount", filter=Q(direction=BankMovement.Direction.OUT)),
    )
    imported_closing = statement.opening_balance + (totals["incoming"] or Decimal("0.00")) - (
        totals["outgoing"] or Decimal("0.00")
    )
    if imported_closing != statement.closing_balance:
        raise TreasuryDomainError("El saldo final no coincide con los movimientos del extracto.")
    if bank_account_balance_at(statement.account, statement.date_to) != statement.closing_balance:
        raise TreasuryDomainError("El saldo final del banco no coincide con el saldo del sistema.")
    statement.status = BankStatement.Status.RECONCILED
    statement.reconciled_by = user
    statement.reconciled_at = timezone.now()
    statement.save(update_fields=("status", "reconciled_by", "reconciled_at", "updated_at"))
    AuditLog.objects.create(
        user=user, action="BANK_STATEMENT_CLOSE", entity="BankStatement",
        entity_id=str(statement.pk), meta_data={"closing_balance": str(statement.closing_balance)},
    )
    return statement


@transaction.atomic
def register_cash_safe_withdrawal(
    *, company_id, store_id, movement_type, amount, currency,
    recipient_name, fund, destination_fund=None, destination_register=None,
    destination="", reference="", notes="",
    bank_account=None,
    supplier=None, purchase_document=None, expense_category=None,
    supporting_document_number="", user=None,
):
    # Serializa retiros de la misma tienda para impedir dos salidas sobre el mismo saldo.
    Company.objects.select_for_update().get(pk=company_id)
    store = Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    if (
        not fund.active or str(fund.company_id) != str(company_id)
        or str(fund.store_id) != str(store_id)
    ):
        raise TreasuryDomainError("Seleccione un fondo activo de la tienda.")
    available = cash_safe_balance(
        company_id=company_id, store_id=store.pk, currency=currency, fund_id=fund.pk
    )
    if amount > available:
        raise TreasuryDomainError(
            f"El importe supera el saldo disponible ({available:.2f} {currency})."
        )
    if movement_type == CashSafeMovement.MovementType.BANK_DEPOSIT:
        if not bank_account or str(bank_account.company_id) != str(company_id):
            raise TreasuryDomainError("Seleccione una cuenta bancaria válida.")
        if bank_account.currency != currency:
            raise TreasuryDomainError("La moneda no coincide con la cuenta bancaria.")
        destination = str(bank_account)
    if movement_type == CashSafeMovement.MovementType.FUND_TRANSFER:
        if (
            not destination_fund or not destination_fund.active
            or str(destination_fund.company_id) != str(company_id)
            or str(destination_fund.store_id) != str(store_id)
        ):
            raise TreasuryDomainError("Seleccione un fondo destino activo de la tienda.")
        if destination_fund.pk == fund.pk:
            raise TreasuryDomainError("El fondo destino debe ser diferente al origen.")
        destination = destination_fund.name
    if movement_type == CashSafeMovement.MovementType.DRAWER_TRANSFER:
        if (
            not destination_register or not destination_register.active
            or str(destination_register.company_id) != str(company_id)
            or str(destination_register.store_id) != str(store_id)
        ):
            raise TreasuryDomainError("Seleccione una caja POS activa de la tienda.")
        destination = f"{destination_register.code} · {destination_register.name}"
    movement = CashSafeMovement(
        company_id=company_id,
        store=store,
        fund=fund,
        destination_fund=destination_fund,
        destination_register=destination_register,
        direction=CashSafeMovement.Direction.OUT,
        movement_type=movement_type,
        amount=amount,
        currency=currency,
        recipient_name=recipient_name,
        destination=destination,
        reference=reference,
        notes=notes,
        bank_account=bank_account,
        supplier=supplier,
        purchase_document=purchase_document,
        expense_category=expense_category,
        supporting_document_number=supporting_document_number,
        created_by=user,
        status=CashSafeMovement.Status.PENDING,
    )
    movement.sequence_number, movement.movement_number = _next_cash_safe_number(company_id)
    movement.full_clean()
    movement.save()
    if movement_type == CashSafeMovement.MovementType.BANK_DEPOSIT:
        CashSafeDeposit.objects.create(
            cash_movement=movement, account=bank_account,
            status=CashSafeDeposit.Status.REQUESTED,
        )
    AuditLog.objects.create(
        user=user,
        action="CASH_SAFE_OUT",
        entity="CashSafeMovement",
        entity_id=str(movement.pk),
        meta_data={
            "amount": str(amount),
            "currency": currency,
            "recipient": recipient_name,
            "destination": destination,
            "movement_type": movement_type,
            "fund_id": str(fund.pk),
            "destination_fund_id": str(destination_fund.pk) if destination_fund else "",
            "destination_register_id": (
                str(destination_register.pk) if destination_register else ""
            ),
        },
    )
    return movement


@transaction.atomic
def register_cash_safe_income(
    *, company_id, store_id, fund, amount, currency, source_name,
    reference="", notes="", user=None,
):
    store = Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    if (
        not fund.active or str(fund.company_id) != str(company_id)
        or str(fund.store_id) != str(store_id)
    ):
        raise TreasuryDomainError("Seleccione un fondo activo de la tienda.")
    movement = CashSafeMovement(
        company_id=company_id, store=store, fund=fund,
        direction=CashSafeMovement.Direction.IN,
        movement_type=CashSafeMovement.MovementType.MANUAL_INCOME,
        amount=amount, currency=currency,
        recipient_name=source_name, destination="Caja fuerte de tienda",
        reference=reference, notes=notes, created_by=user,
        status=CashSafeMovement.Status.PENDING,
    )
    movement.sequence_number, movement.movement_number = _next_cash_safe_number(company_id)
    movement.full_clean()
    movement.save()
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_IN_REQUEST", entity="CashSafeMovement",
        entity_id=str(movement.pk), meta_data={
            "amount": str(amount), "currency": currency, "source": source_name,
            "fund_id": str(fund.pk),
        },
    )
    return movement


@transaction.atomic
def approve_cash_safe_movement(*, movement_id, company_id, store_id, user):
    Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    movement = CashSafeMovement.objects.select_for_update().get(
        pk=movement_id, company_id=company_id, store_id=store_id
    )
    if movement.status != CashSafeMovement.Status.PENDING:
        raise TreasuryDomainError("El movimiento ya fue atendido.")
    if movement.created_by_id == user.pk and user_store_role(
        user, company_id, store_id
    ) != "ADMIN":
        raise TreasuryDomainError("La misma persona no puede atender su propia solicitud.")
    if movement.direction == CashSafeMovement.Direction.OUT:
        available = cash_safe_balance(
            company_id=company_id, store_id=store_id, currency=movement.currency,
            fund_id=movement.fund_id,
        )
        if movement.amount > available:
            raise TreasuryDomainError("El saldo disponible cambió y ya no cubre esta salida.")
    movement.status = CashSafeMovement.Status.POSTED
    movement.authorized_by = user
    movement.authorized_at = timezone.now()
    movement.save(update_fields=("status", "authorized_by", "authorized_at", "updated_at"))
    counterpart = None
    if movement.movement_type == CashSafeMovement.MovementType.FUND_TRANSFER:
        if not movement.destination_fund_id:
            raise TreasuryDomainError("La transferencia no tiene un fondo destino.")
        try:
            destination_fund = TreasuryFund.objects.select_for_update().get(
                pk=movement.destination_fund_id, company_id=company_id,
                store_id=store_id, active=True,
            )
        except TreasuryFund.DoesNotExist as exc:
            raise TreasuryDomainError("El fondo destino ya no está activo.") from exc
        counterpart = CashSafeMovement(
            company_id=company_id, store_id=store_id, fund=destination_fund,
            direction=CashSafeMovement.Direction.IN,
            movement_type=CashSafeMovement.MovementType.FUND_TRANSFER,
            amount=movement.amount, currency=movement.currency,
            recipient_name=movement.recipient_name,
            destination=destination_fund.name,
            reference=movement.movement_number,
            notes=f"Ingreso vinculado desde {movement.fund.name}. {movement.notes}".strip(),
            created_by=movement.created_by, authorized_by=user,
            authorized_at=movement.authorized_at,
            status=CashSafeMovement.Status.POSTED,
            transfer_source=movement,
        )
        counterpart.sequence_number, counterpart.movement_number = _next_cash_safe_number(company_id)
        counterpart.full_clean()
        counterpart.save()
    if movement.movement_type == CashSafeMovement.MovementType.BANK_DEPOSIT:
        try:
            tracking = movement.deposit_tracking
        except CashSafeDeposit.DoesNotExist:
            if not movement.bank_account_id:
                raise TreasuryDomainError("El depósito no tiene una cuenta bancaria de destino.")
            tracking = CashSafeDeposit.objects.create(
                cash_movement=movement, account=movement.bank_account
            )
        tracking.status = CashSafeDeposit.Status.IN_TRANSIT
        tracking.handed_over_by = user
        tracking.handed_over_at = timezone.now()
        tracking.save(update_fields=(
            "status", "handed_over_by", "handed_over_at", "updated_at",
        ))
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_APPROVE", entity="CashSafeMovement",
        entity_id=str(movement.pk), meta_data={
            "movement_number": movement.movement_number,
            "transfer_counterpart_id": str(counterpart.pk) if counterpart else "",
        },
    )
    return movement


@transaction.atomic
def reject_cash_safe_movement(*, movement_id, company_id, store_id, reason, user):
    movement = CashSafeMovement.objects.select_for_update().get(
        pk=movement_id, company_id=company_id, store_id=store_id
    )
    if movement.status != CashSafeMovement.Status.PENDING:
        raise TreasuryDomainError("El movimiento ya fue atendido.")
    if movement.created_by_id == user.pk and user_store_role(
        user, company_id, store_id
    ) != "ADMIN":
        raise TreasuryDomainError("La misma persona no puede atender su propia solicitud.")
    movement.status = CashSafeMovement.Status.REJECTED
    movement.authorized_by = user
    movement.authorized_at = timezone.now()
    movement.reversal_reason = reason
    movement.save(update_fields=(
        "status", "authorized_by", "authorized_at", "reversal_reason", "updated_at",
    ))
    if movement.movement_type == CashSafeMovement.MovementType.BANK_DEPOSIT:
        deposits = CashSafeDeposit.objects.filter(cash_movement=movement)
        deposits.filter(status=CashSafeDeposit.Status.DEPOSITED).update(
            status=CashSafeDeposit.Status.REVERSED
        )
        deposits.exclude(status=CashSafeDeposit.Status.REVERSED).update(
            status=CashSafeDeposit.Status.CANCELLED
        )
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_REJECT", entity="CashSafeMovement",
        entity_id=str(movement.pk), meta_data={"reason": reason},
    )
    return movement


@transaction.atomic
def reverse_cash_safe_movement(*, movement_id, company_id, store_id, reason, user):
    Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    movement = CashSafeMovement.objects.select_for_update().get(
        pk=movement_id, company_id=company_id, store_id=store_id
    )
    if movement.status != CashSafeMovement.Status.POSTED:
        raise TreasuryDomainError("Solo se puede revertir un movimiento confirmado.")
    if movement.transfer_source_id:
        raise TreasuryDomainError(
            "Revierta la transferencia desde su movimiento de salida original."
        )
    transfer_counterpart = None
    if movement.movement_type == CashSafeMovement.MovementType.FUND_TRANSFER:
        try:
            transfer_counterpart = CashSafeMovement.objects.select_for_update().get(
                transfer_source=movement, status=CashSafeMovement.Status.POSTED
            )
        except CashSafeMovement.DoesNotExist as exc:
            raise TreasuryDomainError(
                "No se encontró la contrapartida confirmada de la transferencia."
            ) from exc
        destination_available = cash_safe_balance(
            company_id=company_id, store_id=store_id,
            currency=transfer_counterpart.currency,
            fund_id=transfer_counterpart.fund_id,
        )
        if transfer_counterpart.amount > destination_available:
            raise TreasuryDomainError(
                "El fondo destino ya no tiene saldo suficiente para revertir la transferencia."
            )
    reversal_direction = (
        CashSafeMovement.Direction.OUT
        if movement.direction == CashSafeMovement.Direction.IN
        else CashSafeMovement.Direction.IN
    )
    if reversal_direction == CashSafeMovement.Direction.OUT:
        available = cash_safe_balance(
            company_id=company_id, store_id=store_id, currency=movement.currency,
            fund_id=movement.fund_id,
        )
        if movement.amount > available:
            raise TreasuryDomainError("El saldo actual no permite revertir este ingreso.")
    reversal = CashSafeMovement(
        company_id=company_id, store_id=store_id, fund=movement.fund,
        direction=reversal_direction,
        movement_type=CashSafeMovement.MovementType.ADJUSTMENT,
        amount=movement.amount, currency=movement.currency,
        recipient_name=movement.recipient_name,
        destination="Reversión de " + movement.movement_number,
        reference=movement.reference,
        notes=reason,
        created_by=user, authorized_by=user, authorized_at=timezone.now(),
        status=CashSafeMovement.Status.POSTED,
        reversed_movement=movement,
    )
    reversal.sequence_number, reversal.movement_number = _next_cash_safe_number(company_id)
    reversal.full_clean()
    reversal.save()
    destination_reversal = None
    if transfer_counterpart:
        destination_reversal = CashSafeMovement(
            company_id=company_id, store_id=store_id, fund=transfer_counterpart.fund,
            direction=CashSafeMovement.Direction.OUT,
            movement_type=CashSafeMovement.MovementType.ADJUSTMENT,
            amount=transfer_counterpart.amount, currency=transfer_counterpart.currency,
            recipient_name=transfer_counterpart.recipient_name,
            destination="Reversión de " + transfer_counterpart.movement_number,
            reference=movement.movement_number, notes=reason,
            created_by=user, authorized_by=user, authorized_at=timezone.now(),
            status=CashSafeMovement.Status.POSTED,
            reversed_movement=transfer_counterpart,
        )
        destination_reversal.sequence_number, destination_reversal.movement_number = _next_cash_safe_number(company_id)
        destination_reversal.full_clean()
        destination_reversal.save()
        transfer_counterpart.status = CashSafeMovement.Status.REVERSED
        transfer_counterpart.reversed_at = timezone.now()
        transfer_counterpart.reversal_reason = reason
        transfer_counterpart.save(update_fields=(
            "status", "reversed_at", "reversal_reason", "updated_at",
        ))
    try:
        source_bank_movement = movement.bank_movement
    except BankMovement.DoesNotExist:
        source_bank_movement = None
    if source_bank_movement:
        bank_reversal = BankMovement(
            company_id=company_id,
            account=source_bank_movement.account,
            direction=(
                BankMovement.Direction.OUT
                if source_bank_movement.direction == BankMovement.Direction.IN
                else BankMovement.Direction.IN
            ),
            movement_type=BankMovement.MovementType.ADJUSTMENT,
            amount=source_bank_movement.amount,
            operation_date=timezone.localdate(),
            reference=source_bank_movement.movement_number,
            description=f"Reversión bancaria por {movement.movement_number}",
            safe_movement=reversal,
            reversed_movement=source_bank_movement,
            created_by=user,
        )
        bank_reversal.sequence_number, bank_reversal.movement_number = _next_bank_movement_number(
            company_id, bank_reversal.operation_date
        )
        bank_reversal.full_clean()
        bank_reversal.save()
        source_bank_movement.status = BankMovement.Status.REVERSED
        source_bank_movement.reversed_at = timezone.now()
        source_bank_movement.reversal_reason = reason
        source_bank_movement.save(update_fields=(
            "status", "reversed_at", "reversal_reason", "updated_at",
        ))
    movement.status = CashSafeMovement.Status.REVERSED
    movement.reversed_at = timezone.now()
    movement.reversal_reason = reason
    movement.save(update_fields=("status", "reversed_at", "reversal_reason", "updated_at"))
    if movement.movement_type == CashSafeMovement.MovementType.BANK_DEPOSIT:
        deposits = CashSafeDeposit.objects.filter(cash_movement=movement)
        deposits.filter(status=CashSafeDeposit.Status.DEPOSITED).update(
            status=CashSafeDeposit.Status.REVERSED
        )
        deposits.exclude(status=CashSafeDeposit.Status.REVERSED).update(
            status=CashSafeDeposit.Status.CANCELLED
        )
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_REVERSE", entity="CashSafeMovement",
        entity_id=str(movement.pk),
        meta_data={
            "reason": reason, "reversal_id": str(reversal.pk),
            "destination_reversal_id": (
                str(destination_reversal.pk) if destination_reversal else ""
            ),
        },
    )
    return reversal


@transaction.atomic
def confirm_cash_safe_deposit(
    *, deposit_id, company_id, store_id, operation_date,
    operation_reference, confirmation_notes, user,
):
    Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    deposit = CashSafeDeposit.objects.select_for_update().select_related(
        "cash_movement", "account"
    ).get(
        pk=deposit_id,
        cash_movement__company_id=company_id,
        cash_movement__store_id=store_id,
    )
    if deposit.status != CashSafeDeposit.Status.IN_TRANSIT:
        raise TreasuryDomainError("Solo se puede confirmar un depósito que está en tránsito.")
    movement = deposit.cash_movement
    if movement.status != CashSafeMovement.Status.POSTED:
        raise TreasuryDomainError("La salida de caja fuerte no está confirmada.")
    if not operation_reference:
        raise TreasuryDomainError("Ingrese el número de operación o voucher bancario.")
    if deposit.handed_over_at and operation_date < timezone.localdate(deposit.handed_over_at):
        raise TreasuryDomainError("La fecha del depósito no puede ser anterior a la entrega.")
    bank_movement = BankMovement(
        company_id=company_id, account=deposit.account,
        direction=BankMovement.Direction.IN,
        movement_type=BankMovement.MovementType.SAFE_DEPOSIT,
        amount=movement.amount, operation_date=operation_date,
        reference=operation_reference,
        description=f"Depósito desde caja fuerte {movement.movement_number}",
        safe_movement=movement, created_by=user,
    )
    bank_movement.sequence_number, bank_movement.movement_number = _next_bank_movement_number(
        company_id, operation_date
    )
    bank_movement.full_clean()
    bank_movement.save()
    deposit.status = CashSafeDeposit.Status.DEPOSITED
    deposit.deposited_by = user
    deposit.deposited_at = timezone.now()
    deposit.operation_date = operation_date
    deposit.operation_reference = operation_reference
    deposit.confirmation_notes = confirmation_notes
    deposit.save(update_fields=(
        "status", "deposited_by", "deposited_at", "operation_date",
        "operation_reference", "confirmation_notes", "updated_at",
    ))
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_DEPOSIT_CONFIRM", entity="CashSafeDeposit",
        entity_id=str(deposit.pk), meta_data={
            "cash_movement": movement.movement_number,
            "bank_movement": bank_movement.movement_number,
            "operation_reference": operation_reference,
        },
    )
    return deposit


@transaction.atomic
def register_closing_deposit(*, cash_session, amount, user):
    fund = get_or_create_default_fund(
        company_id=cash_session.company_id, store_id=cash_session.store_id
    )
    movement = CashSafeMovement(
        company=cash_session.company, store=cash_session.store, fund=fund,
        cash_session=cash_session,
        direction=CashSafeMovement.Direction.IN,
        movement_type=CashSafeMovement.MovementType.CLOSING_DEPOSIT,
        amount=amount, currency=cash_session.currency,
        destination="Caja fuerte de tienda",
        notes=f"Cierre de caja N.º {cash_session.session_number}",
        created_by=user, authorized_by=user, authorized_at=timezone.now(),
        status=CashSafeMovement.Status.POSTED,
    )
    movement.sequence_number, movement.movement_number = _next_cash_safe_number(
        cash_session.company_id
    )
    movement.full_clean()
    movement.save()
    return movement


@transaction.atomic
def register_bank_movement(*, company_id, user, **data):
    Company.objects.select_for_update().get(pk=company_id)
    account = BankAccount.objects.select_for_update().get(
        pk=data["account"].pk, company_id=company_id, active=True
    )
    destination_account = data.get("destination_account")
    amount = data["amount"]
    if data["direction"] == BankMovement.Direction.OUT:
        if not account.use_for_payments:
            raise TreasuryDomainError("La cuenta no está habilitada para registrar egresos.")
        if (
            data["movement_type"] == BankMovement.MovementType.TRANSFER
            and account.max_transfer_amount
            and amount > account.max_transfer_amount
        ):
            raise TreasuryDomainError("La transferencia supera el límite configurado para la cuenta.")
        if amount > bank_account_balance(account):
            raise TreasuryDomainError("El egreso supera el saldo contable de la cuenta.")
    if data["movement_type"] == BankMovement.MovementType.TRANSFER:
        if data["direction"] != BankMovement.Direction.OUT:
            raise TreasuryDomainError("Las transferencias internas deben registrarse desde Egresos.")
        if not destination_account:
            raise TreasuryDomainError("Seleccione una cuenta bancaria destino.")
        destination_account = BankAccount.objects.select_for_update().get(
            pk=destination_account.pk, company_id=company_id, active=True
        )
        if destination_account.pk == account.pk:
            raise TreasuryDomainError("La cuenta destino debe ser diferente al origen.")
        if destination_account.currency != account.currency:
            raise TreasuryDomainError("Las cuentas de la transferencia deben usar la misma moneda.")
    movement = BankMovement(company_id=company_id, created_by=user, **data)
    movement.account = account
    movement.sequence_number, movement.movement_number = _next_bank_movement_number(
        company_id, movement.operation_date
    )
    movement.full_clean()
    movement.save()
    counterpart = None
    if movement.movement_type == BankMovement.MovementType.TRANSFER:
        counterpart = BankMovement(
            company_id=company_id,
            account=destination_account,
            direction=BankMovement.Direction.IN,
            movement_type=BankMovement.MovementType.TRANSFER,
            amount=movement.amount,
            operation_date=movement.operation_date,
            reference=movement.movement_number,
            description=f"Transferencia recibida desde {account}",
            transfer_source=movement,
            created_by=user,
        )
        counterpart.sequence_number, counterpart.movement_number = _next_bank_movement_number(
            company_id, counterpart.operation_date
        )
        counterpart.full_clean()
        counterpart.save()
    AuditLog.objects.create(
        user=user, action="BANK_MOVEMENT_CREATE", entity="BankMovement",
        entity_id=str(movement.pk), meta_data={
            "number": movement.movement_number,
            "transfer_counterpart_id": str(counterpart.pk) if counterpart else "",
        },
    )
    return movement


@transaction.atomic
def reverse_bank_movement(*, movement_id, company_id, reason, user):
    if not reason:
        raise TreasuryDomainError("Indique el motivo de la reversión bancaria.")
    Company.objects.select_for_update().get(pk=company_id)
    movement = BankMovement.objects.select_for_update(of=("self",)).select_related(
        "account", "destination_account", "transfer_source"
    ).get(pk=movement_id, company_id=company_id)
    if movement.status != BankMovement.Status.POSTED:
        raise TreasuryDomainError("Solo se puede revertir un movimiento confirmado.")
    if movement.transfer_source_id:
        raise TreasuryDomainError("Revierta la transferencia desde el egreso original.")
    if BankReconciliationMatch.objects.filter(bank_movement=movement).exists():
        raise TreasuryDomainError("No se puede revertir un movimiento ya conciliado.")

    counterpart = None
    if movement.movement_type == BankMovement.MovementType.TRANSFER:
        try:
            counterpart = BankMovement.objects.select_for_update().select_related("account").get(
                transfer_source=movement, status=BankMovement.Status.POSTED
            )
        except BankMovement.DoesNotExist as exc:
            raise TreasuryDomainError("No se encontró la contrapartida de la transferencia.") from exc
        if BankReconciliationMatch.objects.filter(bank_movement=counterpart).exists():
            raise TreasuryDomainError("La contrapartida de la transferencia ya está conciliada.")
        if counterpart.amount > bank_account_balance(counterpart.account):
            raise TreasuryDomainError("La cuenta destino ya no tiene saldo para revertir la transferencia.")
    elif (
        movement.direction == BankMovement.Direction.IN
        and movement.amount > bank_account_balance(movement.account)
    ):
        raise TreasuryDomainError("La cuenta ya no tiene saldo para revertir este ingreso.")

    now = timezone.now()
    reversal = BankMovement(
        company_id=company_id,
        account=movement.account,
        direction=(
            BankMovement.Direction.OUT
            if movement.direction == BankMovement.Direction.IN
            else BankMovement.Direction.IN
        ),
        movement_type=BankMovement.MovementType.ADJUSTMENT,
        amount=movement.amount,
        operation_date=timezone.localdate(),
        reference=movement.movement_number,
        description=f"Reversión de {movement.movement_number}: {reason}",
        reversed_movement=movement,
        created_by=user,
    )
    reversal.sequence_number, reversal.movement_number = _next_bank_movement_number(
        company_id, reversal.operation_date
    )
    reversal.full_clean()
    reversal.save()

    destination_reversal = None
    if counterpart:
        destination_reversal = BankMovement(
            company_id=company_id,
            account=counterpart.account,
            direction=BankMovement.Direction.OUT,
            movement_type=BankMovement.MovementType.ADJUSTMENT,
            amount=counterpart.amount,
            operation_date=timezone.localdate(),
            reference=counterpart.movement_number,
            description=f"Reversión de {counterpart.movement_number}: {reason}",
            reversed_movement=counterpart,
            created_by=user,
        )
        (
            destination_reversal.sequence_number,
            destination_reversal.movement_number,
        ) = _next_bank_movement_number(company_id, destination_reversal.operation_date)
        destination_reversal.full_clean()
        destination_reversal.save()
        counterpart.status = BankMovement.Status.REVERSED
        counterpart.reversed_at = now
        counterpart.reversal_reason = reason
        counterpart.save(update_fields=(
            "status", "reversed_at", "reversal_reason", "updated_at",
        ))

    movement.status = BankMovement.Status.REVERSED
    movement.reversed_at = now
    movement.reversal_reason = reason
    movement.save(update_fields=(
        "status", "reversed_at", "reversal_reason", "updated_at",
    ))
    AuditLog.objects.create(
        user=user, action="BANK_MOVEMENT_REVERSE", entity="BankMovement",
        entity_id=str(movement.pk), meta_data={
            "reason": reason,
            "reversal_id": str(reversal.pk),
            "destination_reversal_id": (
                str(destination_reversal.pk) if destination_reversal else ""
            ),
        },
    )
    return reversal


def register_linked_collection(
    *, account, amount, operation_date, reference, description,
    source_type, source_id, user, direction=BankMovement.Direction.IN,
):
    if account.currency is None:
        raise TreasuryDomainError("La cuenta bancaria no tiene moneda configurada.")
    return register_bank_movement(
        company_id=account.company_id,
        user=user,
        account=account,
        direction=direction,
        movement_type=(
            BankMovement.MovementType.RECEIPT
            if direction == BankMovement.Direction.IN
            else BankMovement.MovementType.PAYMENT
        ),
        amount=amount,
        operation_date=operation_date,
        reference=reference,
        description=description,
        source_type=source_type,
        source_id=str(source_id),
    )


def _cash_safe_count_number(counted_at, sequence_number):
    local_counted_at = timezone.localtime(counted_at) if timezone.is_aware(counted_at) else counted_at
    return f"ARQ-{local_counted_at:%Y%m}-{sequence_number:04d}"


@transaction.atomic
def register_cash_safe_count(
    *, company_id, store_id, counted_at, currency, counted_amount,
    fund, denominations=None, notes="", user=None,
):
    from .models import CashSafeCount

    Company.objects.select_for_update().get(pk=company_id)
    store = Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    if (
        not fund.active or str(fund.company_id) != str(company_id)
        or str(fund.store_id) != str(store_id)
    ):
        raise TreasuryDomainError("Seleccione un fondo activo de la tienda.")
    now = timezone.now()
    if counted_at.second == 0 and counted_at.microsecond == 0 and (
        timezone.localtime(counted_at).replace(second=0, microsecond=0)
        == timezone.localtime(now).replace(second=0, microsecond=0)
    ):
        counted_at = now
    if counted_at > now:
        raise TreasuryDomainError("La fecha del conteo no puede estar en el futuro.")
    previous_closed = CashSafeCount.objects.filter(
        company_id=company_id, store_id=store_id, currency=currency,
        fund=fund, status=CashSafeCount.Status.CLOSED,
    ).order_by("-counted_at", "-created_at").first()
    previous_boundary = (
        previous_closed.reviewed_at or previous_closed.counted_at
        if previous_closed else None
    )
    if previous_closed and counted_at <= previous_boundary:
        raise TreasuryDomainError(
            f"La fecha debe ser posterior al cierre del último arqueo ({previous_closed.count_number})."
        )
    if CashSafeCount.objects.filter(
        company_id=company_id, store_id=store_id, currency=currency,
        fund=fund,
        status=CashSafeCount.Status.PENDING,
    ).exists():
        raise TreasuryDomainError(
            "Ya existe un arqueo pendiente para este fondo y moneda."
        )
    local_counted_at = timezone.localtime(counted_at) if timezone.is_aware(counted_at) else counted_at
    last_sequence = CashSafeCount.objects.filter(
        company_id=company_id,
        counted_at__year=local_counted_at.year,
        counted_at__month=local_counted_at.month,
    ).aggregate(value=Max("sequence_number"))["value"] or 0
    sequence_number = last_sequence + 1
    theoretical = cash_safe_balance(
        company_id=company_id, store_id=store.pk, currency=currency,
        fund_id=fund.pk, at=counted_at,
    )
    count = CashSafeCount(
        company_id=company_id, store=store, fund=fund, counted_at=counted_at,
        currency=currency, theoretical_amount=theoretical,
        counted_amount=counted_amount, difference=counted_amount - theoretical,
        denominations=denominations or {}, notes=notes, counted_by=user,
        sequence_number=sequence_number,
    )
    count.count_number = _cash_safe_count_number(counted_at, sequence_number)
    count.full_clean()
    count.save()
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_COUNT_CREATE", entity="CashSafeCount",
        entity_id=str(count.pk), meta_data={
            "number": count.count_number, "theoretical": str(theoretical),
            "counted": str(counted_amount), "difference": str(count.difference),
            "fund_id": str(fund.pk),
        },
    )
    return count


@transaction.atomic
def review_cash_safe_count(
    *, count_id, company_id, store_id, action, review_notes, user,
):
    from .models import CashSafeCount

    Store.objects.select_for_update().get(pk=store_id, company_id=company_id)
    count = CashSafeCount.objects.select_for_update().get(
        pk=count_id, company_id=company_id, store_id=store_id
    )
    if count.status != CashSafeCount.Status.PENDING:
        raise TreasuryDomainError("El arqueo ya fue atendido.")
    if count.counted_by_id == user.pk and user_store_role(user, company_id, store_id) != "ADMIN":
        raise TreasuryDomainError("La misma persona no puede revisar su propio arqueo.")
    if action == "reject":
        if not review_notes:
            raise TreasuryDomainError("Indique el motivo del rechazo.")
        count.status = CashSafeCount.Status.REJECTED
    elif action == "close":
        count.status = CashSafeCount.Status.CLOSED
        if count.difference:
            movement = CashSafeMovement(
                company_id=company_id, store_id=store_id, fund=count.fund,
                direction=(CashSafeMovement.Direction.IN if count.difference > 0
                           else CashSafeMovement.Direction.OUT),
                movement_type=CashSafeMovement.MovementType.ADJUSTMENT,
                amount=abs(count.difference), currency=count.currency,
                destination=f"Ajuste por arqueo {count.count_number}",
                notes=review_notes or "Diferencia confirmada durante el arqueo.",
                created_by=user, authorized_by=user, authorized_at=timezone.now(),
                status=CashSafeMovement.Status.POSTED,
            )
            movement.sequence_number, movement.movement_number = _next_cash_safe_number(company_id)
            movement.full_clean()
            movement.save()
            count.adjustment_movement = movement
    else:
        raise TreasuryDomainError("Acción de revisión inválida.")
    count.reviewed_by = user
    count.reviewed_at = timezone.now()
    count.review_notes = review_notes
    count.save(update_fields=(
        "status", "reviewed_by", "reviewed_at", "review_notes",
        "adjustment_movement", "updated_at",
    ))
    AuditLog.objects.create(
        user=user, action="CASH_SAFE_COUNT_REVIEW", entity="CashSafeCount",
        entity_id=str(count.pk), meta_data={"status": count.status},
    )
    return count
