import io
from collections import defaultdict
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe_text(value):
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def build_cash_session_report(sessions, *, store_name, date_from=None, date_to=None):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    detail = workbook.create_sheet("Cierres de caja")
    header_fill = PatternFill("solid", fgColor="1A7F64")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16, color="123B34")

    period = "Todos los registros"
    if date_from or date_to:
        period = f"Del {date_from:%d/%m/%Y}" if date_from else "Desde el inicio"
        period += f" al {date_to:%d/%m/%Y}" if date_to else " en adelante"
    summary.append(["Reporte consolidado de cierres de caja"])
    summary["A1"].font = title_font
    summary.append(["Sucursal", _safe_text(store_name)])
    summary.append(["Periodo (fecha de apertura)", period])
    summary.append(["Generado", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    summary.append([])
    summary.append([
        "Moneda", "Turnos", "Fondo inicial", "Ventas/cobros", "Esperado",
        "Contado", "Gaveta siguiente", "Caja fuerte", "Banco", "Diferencia",
    ])
    for cell in summary[6]:
        cell.fill = header_fill
        cell.font = header_font

    totals = defaultdict(lambda: {
        "sessions": 0, "opening": Decimal("0"), "collected": Decimal("0"),
        "expected": Decimal("0"), "counted": Decimal("0"),
        "next": Decimal("0"), "safe": Decimal("0"), "bank": Decimal("0"),
        "difference": Decimal("0"),
    })
    rows = []
    for session in sessions:
        currency_totals = totals[session.currency]
        currency_totals["sessions"] += 1
        for key, value in (
            ("opening", session.opening_total),
            ("collected", session.net_collected_total),
            ("expected", session.display_expected_cash_total),
            ("counted", session.counted_cash_total),
            ("next", session.next_opening_total),
            ("safe", session.safe_deposit_total),
            ("bank", session.bank_deposit_total),
            ("difference", session.cash_difference),
        ):
            currency_totals[key] += value or Decimal("0")
        rows.append(session)
    for currency, item in sorted(totals.items()):
        summary.append([
            currency, item["sessions"], float(item["opening"]),
            float(item["collected"]), float(item["expected"]),
            float(item["counted"]), float(item["next"]), float(item["safe"]),
            float(item["bank"]), float(item["difference"]),
        ])

    headers = [
        "N.º turno", "Estado", "Caja", "Sucursal", "Cajero", "Apertura",
        "Cierre", "Moneda", "Fondo inicial", "Ventas/cobros", "Esperado",
        "Contado", "Gaveta siguiente", "Caja fuerte", "Banco", "Diferencia",
        "Cerrado por", "Observación de cierre",
    ]
    detail.append(headers)
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for session in rows:
        detail.append([
            session.session_number, session.get_status_display(), session.register.code,
            _safe_text(session.store.name),
            _safe_text(session.opened_by.display_name if session.opened_by else "Sistema"),
            timezone.localtime(session.opened_at).replace(tzinfo=None),
            timezone.localtime(session.closed_at).replace(tzinfo=None) if session.closed_at else None,
            session.currency, float(session.opening_total),
            float(session.net_collected_total), float(session.display_expected_cash_total),
            float(session.counted_cash_total) if session.counted_cash_total is not None else None,
            float(session.next_opening_total) if session.next_opening_total is not None else None,
            float(session.safe_deposit_total), float(session.bank_deposit_total),
            float(session.cash_difference) if session.cash_difference is not None else None,
            _safe_text(session.closed_by.display_name if session.closed_by else ""),
            _safe_text(session.closing_note),
        ])
    for worksheet in (summary, detail):
        worksheet.freeze_panes = "A7" if worksheet is summary else "A2"
        if worksheet.max_row > 1:
            worksheet.auto_filter.ref = worksheet.dimensions if worksheet is detail else f"A6:J{worksheet.max_row}"
        for column in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column)
            width = max(len(str(cell.value or "")) for cell in worksheet[letter]) + 3
            worksheet.column_dimensions[letter].width = min(max(width, 12), 42)
    for row in detail.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = "dd/mm/yyyy hh:mm"
    for row in summary.iter_rows(min_row=7, min_col=3, max_col=10):
        for cell in row:
            cell.number_format = "#,##0.00"
    for row in detail.iter_rows(min_row=2, min_col=9, max_col=16):
        for cell in row:
            cell.number_format = "#,##0.00"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
